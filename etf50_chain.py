#!/usr/bin/env python3
"""Clean the supplied SSE 50 ETF option daily workbook and query option chains.

Uses only Python's standard library and openpyxl for reading the source XLSX.
No price, Greek, or contract-multiplier values are imputed.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import sqlite3
import statistics
from collections import Counter, deque
from datetime import datetime
from pathlib import Path


SOURCE_COLUMNS = (
    "security_id", "symbol", "trade_date", "call_put", "open", "high",
    "low", "close", "volume", "amount", "open_interest",
    "pre_settle_price", "settle_price", "list_date", "exercise_price",
    "remaining_time", "last_edate", "delta", "gamma", "rho", "theta",
    "vega", "implc_volatlty", "fund_open", "fund_high", "fund_low",
    "fund_close", "fund_volume", "fund_amount", "ten_year",
)

REFERENCE_COLUMNS = (
    "date", "ETF_close", "thscode", "oi", "d_oi", "settle", "iv",
    "delta", "gamma", "theta", "vega", "rho", "delta_exchg",
    "gamma_exchg", "theta_exchg", "vega_exchg", "rho_exchg", "name",
    "hv30", "hv60", "hv90", "mul", "DTE_n", "DTE_t",
)

OPTION_COLUMNS = (
    "date", "thscode", "name", "option_type", "expiry", "strike",
    "list_date", "source_remaining_time", "open", "high", "low", "close",
    "volume", "amount", "oi", "d_oi", "pre_settle_price", "settle",
    "delta", "gamma", "rho", "theta", "vega", "iv", "source_row",
)

UNDERLYING_COLUMNS = (
    "fund_open", "fund_high", "fund_low", "fund_close", "fund_volume",
    "fund_amount", "ten_year",
)

FULL_COLUMNS = (
    *OPTION_COLUMNS, *UNDERLYING_COLUMNS, "hv30", "hv60", "hv90",
    "DTE_n", "DTE_t",
)


def _date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, int):
        value = str(value)
    if not isinstance(value, str):
        raise ValueError(f"Invalid source date: {value!r}")
    value = value.strip()
    if len(value) == 8 and value.isdigit():
        return datetime.strptime(value, "%Y%m%d").date().isoformat()
    return datetime.strptime(value, "%Y-%m-%d").date().isoformat()


def _number(value: object, *, integer: bool = False) -> float | int | None:
    if value is None or value == "":
        return None
    result = float(value)
    if not math.isfinite(result):
        raise ValueError(f"Non-finite numeric value: {value!r}")
    if integer:
        if not result.is_integer():
            raise ValueError(f"Expected whole number: {value!r}")
        return int(result)
    return result


def _option_row(raw: tuple, row_no: int) -> tuple:
    if not raw[0] or not raw[1]:
        raise ValueError(f"Missing contract ID or name at row {row_no}")
    date, expiry, listed = (_date(raw[i]) for i in (2, 16, 13))
    if raw[3] not in ("C", "P"):
        raise ValueError(f"Unknown call/put value at row {row_no}: {raw[3]!r}")
    dte = (datetime.fromisoformat(expiry) - datetime.fromisoformat(date)).days
    if dte < 0 or _number(raw[15], integer=True) != dte:
        raise ValueError(f"Inconsistent expiry/remaining_time at row {row_no}")
    return (
        date, str(raw[0]), str(raw[1]), raw[3], expiry,
        _number(raw[14]), listed, dte,
        *(_number(raw[i]) for i in (4, 5, 6, 7)),
        _number(raw[8], integer=True), _number(raw[9]),
        _number(raw[10], integer=True), None,  # d_oi is set after deduplication
        _number(raw[11]), _number(raw[12]),
        *(_number(raw[i]) for i in (17, 18, 19, 20, 21, 22)),
        row_no,
    )


def _underlying_row(raw: tuple) -> tuple:
    return tuple(
        _number(raw[i], integer=i == 27) for i in range(23, 30)
    )


def _create_schema(db: sqlite3.Connection) -> None:
    db.executescript("""
        CREATE TABLE option_daily (
            date TEXT NOT NULL,
            thscode TEXT NOT NULL,
            name TEXT NOT NULL,
            option_type TEXT NOT NULL CHECK(option_type IN ('C','P')),
            expiry TEXT NOT NULL,
            strike REAL,
            list_date TEXT NOT NULL,
            source_remaining_time INTEGER NOT NULL,
            open REAL, high REAL, low REAL, close REAL,
            volume INTEGER, amount REAL, oi INTEGER, d_oi INTEGER,
            pre_settle_price REAL, settle REAL,
            delta REAL, gamma REAL, rho REAL, theta REAL, vega REAL, iv REAL,
            source_row INTEGER NOT NULL,
            PRIMARY KEY (date, thscode)
        );
        CREATE INDEX idx_option_chain
            ON option_daily(date, expiry, strike, option_type);

        CREATE TABLE underlying_daily (
            date TEXT PRIMARY KEY,
            fund_open REAL, fund_high REAL, fund_low REAL, fund_close REAL,
            fund_volume INTEGER, fund_amount REAL, ten_year REAL,
            hv30 REAL, hv60 REAL, hv90 REAL
        );
        CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
    """)


def _add_volatility(db: sqlite3.Connection) -> None:
    """Rolling sample SD of N log returns * sqrt(252) * 100 (percent)."""
    windows = {n: deque(maxlen=n) for n in (30, 60, 90)}
    previous_close = None
    updates = []
    for date, close in db.execute(
        "SELECT date, fund_close FROM underlying_daily ORDER BY date"
    ):
        if previous_close is None or close is None or previous_close <= 0 or close <= 0:
            for window in windows.values():
                window.clear()
        else:
            ret = math.log(close / previous_close)
            for window in windows.values():
                window.append(ret)
        vals = [
            statistics.stdev(window) * math.sqrt(252) * 100
            if len(window) == n else None
            for n, window in windows.items()
        ]
        updates.append((*vals, date))
        previous_close = close
    db.executemany(
        "UPDATE underlying_daily SET hv30=?, hv60=?, hv90=? WHERE date=?",
        updates,
    )


def build_database(source: str | Path, db_path: str | Path, *, overwrite: bool = False) -> dict:
    """Ingest every source row, dropping only identical duplicate contract-days."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("To build the database, install openpyxl") from exc

    source, db_path = Path(source), Path(db_path)
    if not source.is_file():
        raise FileNotFoundError(source)
    if db_path.exists():
        if not overwrite:
            raise FileExistsError(f"Database exists: {db_path}. Pass --overwrite to rebuild.")
        db_path.unlink()
    db_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError(f"Expected one source worksheet, got {workbook.sheetnames}")
    rows = iter(workbook.active.iter_rows(values_only=True))
    header = tuple(next(rows))
    if header != SOURCE_COLUMNS:
        raise ValueError(f"Unexpected source columns: {header!r}")

    counts: Counter = Counter()
    missing: Counter = Counter()
    last_oi: dict[str, int | None] = {}
    fund_by_date: dict[str, tuple] = {}
    first_date, last_date = None, None
    db = sqlite3.connect(db_path)
    try:
        _create_schema(db)
        db.execute("BEGIN")
        insert_sql = (
            "INSERT INTO option_daily (" + ",".join(OPTION_COLUMNS) + ") VALUES ("
            + ",".join("?" for _ in OPTION_COLUMNS) + ")"
        )
        for row_no, raw in enumerate(rows, start=2):
            counts["source_rows"] += 1
            if len(raw) != len(SOURCE_COLUMNS):
                raise ValueError(f"Unexpected column count at row {row_no}")
            for column, value in zip(SOURCE_COLUMNS, raw):
                if value is None or value == "":
                    missing[column] += 1
            option = list(_option_row(raw, row_no))
            date, code = option[:2]
            if last_date is not None and date < last_date:
                raise ValueError(f"Source dates are not sorted at row {row_no}")
            first_date = first_date or date
            last_date = date
            fund = _underlying_row(raw)
            if date in fund_by_date and fund_by_date[date] != fund:
                raise ValueError(f"Conflicting underlying values on {date}, row {row_no}")
            fund_by_date[date] = fund

            oi = option[OPTION_COLUMNS.index("oi")]
            prior_oi = last_oi.get(code, 0)
            option[OPTION_COLUMNS.index("d_oi")] = (
                oi - prior_oi if oi is not None and prior_oi is not None else None
            )
            try:
                db.execute(insert_sql, option)
            except sqlite3.IntegrityError as exc:
                prior = db.execute(
                    "SELECT " + ",".join(OPTION_COLUMNS) +
                    " FROM option_daily WHERE date=? AND thscode=?", (date, code)
                ).fetchone()
                # The second identical source row has a different source_row,
                # and its tentative d_oi uses the just-inserted row as prior.
                # Compare only the source-backed fields.
                source_backed = [i for i, name in enumerate(OPTION_COLUMNS)
                                 if name not in ("d_oi", "source_row")]
                if prior is None or any(prior[i] != option[i] for i in source_backed):
                    raise ValueError(
                        f"Conflicting duplicate contract/day at source row {row_no}: "
                        f"{date} {code}"
                    ) from exc
                counts["identical_duplicates_removed"] += 1
                continue
            last_oi[code] = oi
            counts["clean_rows"] += 1

        db.executemany(
            "INSERT INTO underlying_daily (date," + ",".join(UNDERLYING_COLUMNS) +
            ") VALUES (?,?,?,?,?,?,?,?)",
            ((date, *fund) for date, fund in fund_by_date.items()),
        )
        _add_volatility(db)
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        report = {
            "source_file": source.name, "source_sha256": digest,
            "source_rows": counts["source_rows"], "clean_rows": counts["clean_rows"],
            "identical_duplicates_removed": counts["identical_duplicates_removed"],
            "trading_days": len(fund_by_date), "first_date": first_date,
            "last_date": last_date, "missing_source_values": dict(missing),
        }
        db.executemany(
            "INSERT INTO metadata (key,value) VALUES (?,?)",
            ((key, json.dumps(value, ensure_ascii=False)) for key, value in report.items()),
        )
        db.commit()
        return report
    except Exception:
        db.rollback()
        db.close()
        db_path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
        if db:
            db.close()


def _normalize_expiry(expiry: str | None) -> tuple[str, ...] | None:
    if expiry is None:
        return None
    if len(expiry) == 7:
        datetime.strptime(expiry, "%Y-%m")
        return ("o.expiry >= ? AND o.expiry < ?", expiry + "-01",
                f"{int(expiry[:4]) + 1}-01-01" if expiry[5:] == "12"
                else f"{expiry[:5]}{int(expiry[5:]) + 1:02d}-01")
    return ("o.expiry = ?", _date(expiry))


def get_available_dates(db_path: str | Path) -> list[str]:
    with sqlite3.connect(db_path) as db:
        return [r[0] for r in db.execute("SELECT date FROM underlying_daily ORDER BY date")]


def get_expiries(db_path: str | Path, trade_date: str) -> list[str]:
    with sqlite3.connect(db_path) as db:
        return [r[0] for r in db.execute(
            "SELECT DISTINCT expiry FROM option_daily WHERE date=? ORDER BY expiry",
            (_date(trade_date),),
        )]


def load_chain(
    db_path: str | Path, trade_date: str, expiry: str | None = None,
    *, schema: str = "reference", option_type: str | None = None,
) -> list[dict]:
    """Return all maturities, a YYYY-MM month, or one YYYY-MM-DD expiry.

    schema='reference' reproduces the sample's 24-column header; unavailable
    exchange Greeks and multiplier are None. schema='full' retains raw quotes.
    """
    if schema not in ("reference", "full"):
        raise ValueError("schema must be 'reference' or 'full'")
    if option_type not in (None, "C", "P"):
        raise ValueError("option_type must be C or P")
    date = _date(trade_date)
    where, params = ["o.date=?"], [date]
    filter_expiry = _normalize_expiry(expiry)
    if filter_expiry:
        where.append(filter_expiry[0])
        params.extend(filter_expiry[1:])
    if option_type:
        where.append("o.option_type=?")
        params.append(option_type)
    # The observed source calendar is complete only through its latest date.
    # Do not estimate future trading days beyond its coverage.
    dte_trading = """CASE WHEN o.expiry <= (SELECT MAX(date) FROM underlying_daily)
        THEN (SELECT COUNT(*) FROM underlying_daily t
              WHERE t.date BETWEEN o.date AND o.expiry)
        ELSE NULL END AS DTE_t"""
    if schema == "reference":
        select = """o.date AS date, u.fund_close AS ETF_close,
            o.thscode AS thscode, o.oi AS oi, o.d_oi AS d_oi,
            o.settle AS settle, o.iv AS iv,
            o.delta AS delta, o.gamma AS gamma, o.theta AS theta,
            o.vega AS vega, o.rho AS rho,
            NULL AS delta_exchg, NULL AS gamma_exchg,
            NULL AS theta_exchg, NULL AS vega_exchg, NULL AS rho_exchg,
            o.name AS name, u.hv30 AS hv30, u.hv60 AS hv60,
            u.hv90 AS hv90, NULL AS mul,
            o.source_remaining_time + 1 AS DTE_n, """ + dte_trading
        columns = REFERENCE_COLUMNS
    else:
        select = (
            ",".join("o." + name for name in OPTION_COLUMNS)
            + "," + ",".join("u." + name for name in UNDERLYING_COLUMNS)
            + ",u.hv30,u.hv60,u.hv90,"
            + "o.source_remaining_time + 1 AS DTE_n," + dte_trading
        )
        columns = FULL_COLUMNS
    sql = (
        "SELECT " + select + " FROM option_daily o "
        "JOIN underlying_daily u ON u.date=o.date WHERE " + " AND ".join(where)
        + " ORDER BY o.expiry, o.strike, o.option_type, o.thscode"
    )
    with sqlite3.connect(db_path) as db:
        rows = db.execute(sql, params).fetchall()
    return [dict(zip(columns, row)) for row in rows]


def export_csv(
    db_path: str | Path, trade_date: str, output: str | Path,
    expiry: str | None = None, *, schema: str = "reference",
    option_type: str | None = None,
) -> int:
    rows = load_chain(db_path, trade_date, expiry, schema=schema,
                      option_type=option_type)
    if not rows:
        raise ValueError("No option quotes found for the selected date and expiry")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=REFERENCE_COLUMNS if schema == "reference" else FULL_COLUMNS,
        )
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    build = commands.add_parser("build", help="Clean XLSX and build SQLite")
    build.add_argument("--input", required=True, type=Path)
    build.add_argument("--db", required=True, type=Path)
    build.add_argument("--overwrite", action="store_true")
    export = commands.add_parser("export", help="Export an option chain")
    export.add_argument("--db", required=True, type=Path)
    export.add_argument("--date", required=True)
    export.add_argument("--expiry", help="YYYY-MM month or YYYY-MM-DD date")
    export.add_argument("--schema", choices=("reference", "full"), default="reference")
    export.add_argument("--option-type", choices=("C", "P"))
    export.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if args.command == "build":
        print(json.dumps(build_database(args.input, args.db, overwrite=args.overwrite),
                         ensure_ascii=False, indent=2))
    else:
        n = export_csv(args.db, args.date, args.output, args.expiry,
                       schema=args.schema, option_type=args.option_type)
        print(f"Exported {n} contracts to {args.output}")


if __name__ == "__main__":
    main()
